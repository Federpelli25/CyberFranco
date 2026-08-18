import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.add_participant_ids import add_participant_ids


class AddParticipantIdsTests(unittest.TestCase):

    def test_creates_copy_with_text_ids_without_overwriting_source(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "legacy.xlsx"
            destination = Path(directory) / "with_ids.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Nome", "Cognome", "Squadra"])
            sheet.append(["Mario", "Rossi", "ROSSI"])
            sheet.append(["Luca", "Bianchi", "BLU"])
            workbook.save(source)
            workbook.close()

            count = add_participant_ids(source, destination)

            self.assertEqual(count, 2)
            original = load_workbook(source, read_only=True)
            migrated = load_workbook(destination, read_only=True)
            try:
                self.assertEqual(original.active.cell(1, 1).value, "Nome")
                self.assertEqual(migrated.active.cell(1, 1).value, "ID")
                self.assertEqual(migrated.active.cell(2, 1).value, "001")
                self.assertEqual(migrated.active.cell(3, 1).value, "002")
            finally:
                original.close()
                migrated.close()

    def test_existing_destination_requires_force(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "legacy.xlsx"
            destination = Path(directory) / "existing.xlsx"
            source.touch()
            destination.write_text("keep", encoding="utf-8")
            with self.assertRaises(FileExistsError):
                add_participant_ids(source, destination)
            self.assertEqual(destination.read_text(encoding="utf-8"), "keep")


if __name__ == "__main__":
    unittest.main()
