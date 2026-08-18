"""Crea una copia di un Excel partecipanti legacy aggiungendo ID sequenziali."""

import argparse
import shutil
from pathlib import Path

from openpyxl import load_workbook


def add_participant_ids(source: Path, destination: Path, force: bool = False) -> int:
    if not source.is_file():
        raise FileNotFoundError(f"File sorgente non trovato: {source}")
    if source.resolve() == destination.resolve():
        raise ValueError("Il file di destinazione deve essere diverso dall'originale.")
    if destination.exists() and not force:
        raise FileExistsError(
            f"Il file esiste già: {destination}. Usa --force per sovrascriverlo."
        )

    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    workbook = load_workbook(destination)
    try:
        sheet = workbook.active
        headers = [
            str(cell.value).strip() if cell.value is not None else ""
            for cell in sheet[1]
        ]
        if "ID" in headers:
            raise ValueError("Il file contiene già la colonna ID.")
        sheet.insert_cols(1)
        sheet.cell(1, 1, "ID")
        participant_count = max(0, sheet.max_row - 1)
        width = max(3, len(str(participant_count)))
        for index, row_number in enumerate(range(2, sheet.max_row + 1), start=1):
            cell = sheet.cell(row_number, 1, f"{index:0{width}d}")
            cell.number_format = "@"
        workbook.save(destination)
    except Exception:
        workbook.close()
        destination.unlink(missing_ok=True)
        raise
    workbook.close()
    return participant_count


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Excel legacy senza colonna ID")
    parser.add_argument("output", nargs="?", type=Path)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    destination = args.output or args.source.with_name(
        f"{args.source.stem}_con_id{args.source.suffix}"
    )
    count = add_participant_ids(args.source, destination, args.force)
    print(f"Creato: {destination} ({count} partecipanti)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
