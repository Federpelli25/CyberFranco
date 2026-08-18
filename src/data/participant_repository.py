import logging
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from src.core.exceptions import ParticipantDataError


logger = logging.getLogger(__name__)


class ParticipantRepository:

    REQUIRED_COLUMNS = {"ID", "Nome", "Cognome", "Squadra"}
    DEFAULT_REAL_FILE = Path("data/partecipanti.xlsx")
    DEFAULT_EXAMPLE_FILE = Path("data/partecipanti_example.xlsx")

    def __init__(self, file_path: str | Path | None = None):
        self.file_path = (
            self._resolve_default_file()
            if file_path is None
            else Path(file_path)
        )
        self.participants = []

    def load(self) -> list[dict]:
        logger.info("Loading participants file: %s", self.file_path)

        if not self.file_path.exists():
            logger.error("Participants file not found: %s", self.file_path)
            raise ParticipantDataError(
                "FILE PARTECIPANTI NON TROVATO\n\n"
                f"Percorso: {self.file_path}"
            )

        try:
            workbook = load_workbook(
                self.file_path,
                read_only=True,
                data_only=True,
            )
        except (OSError, InvalidFileException, ValueError) as exc:
            logger.exception("Participants workbook cannot be read")
            raise ParticipantDataError(
                "FILE PARTECIPANTI NON VALIDO\n\n"
                "Il file non è leggibile oppure è danneggiato."
            ) from exc

        try:
            rows = workbook.active.iter_rows(values_only=True)

            try:
                headers = next(rows)
            except StopIteration as exc:
                logger.error("Participants workbook is empty: %s", self.file_path)
                raise ParticipantDataError(
                    "FILE PARTECIPANTI NON VALIDO\n\n"
                    "Il file Excel è vuoto."
                ) from exc

            headers = [
                str(header).strip() if header is not None else ""
                for header in headers
            ]
            missing_columns = self.REQUIRED_COLUMNS - set(headers)

            if missing_columns:
                missing_text = ", ".join(sorted(missing_columns))
                logger.error(
                    "Participants workbook missing columns: %s",
                    missing_text,
                )
                raise ParticipantDataError(
                    "FILE PARTECIPANTI NON VALIDO\n\n"
                    f"Colonne mancanti: {missing_text}."
                )

            indexes = {
                column: headers.index(column)
                for column in self.REQUIRED_COLUMNS
            }
            participants = []
            seen_ids: dict[str, int] = {}
            name_rows: dict[str, list[int]] = {}

            for row_number, row in enumerate(rows, start=2):
                values = {
                    column: row[indexes[column]]
                    if indexes[column] < len(row)
                    else None
                    for column in self.REQUIRED_COLUMNS
                }

                if not any(values.values()):
                    continue

                if any(value is None or not str(value).strip() for value in values.values()):
                    logger.error(
                        "Incomplete participant row: file=%s row=%d",
                        self.file_path,
                        row_number,
                    )
                    raise ParticipantDataError(
                        "FILE PARTECIPANTI NON VALIDO\n\n"
                        f"La riga {row_number} contiene ID, nome, cognome "
                        "o squadra vuoti."
                    )

                participant = {
                    "id": str(values["ID"]).strip(),
                    "nome": str(values["Nome"]).strip(),
                    "cognome": str(values["Cognome"]).strip(),
                    "squadra": str(values["Squadra"]).strip(),
                }
                participant["nome_completo"] = (
                    f"{participant['nome']} {participant['cognome']}"
                )
                participant["search_name"] = (
                    participant["nome_completo"].casefold()
                )

                if participant["id"] in seen_ids:
                    logger.error(
                        "Duplicate participant ID: file=%s rows=%d,%d id=%s",
                        self.file_path,
                        seen_ids[participant["id"]],
                        row_number,
                        participant["id"],
                    )
                    raise ParticipantDataError(
                        "ID PARTECIPANTE DUPLICATO\n\n"
                        f"L'ID {participant['id']} è presente più volte "
                        "nel file Excel."
                    )

                seen_ids[participant["id"]] = row_number
                name_rows.setdefault(participant["search_name"], []).append(
                    row_number
                )
                participants.append(participant)

            if not participants:
                logger.error("Participants workbook has no valid data rows")
                raise ParticipantDataError(
                    "FILE PARTECIPANTI NON VALIDO\n\n"
                    "Non sono presenti partecipanti validi."
                )

            self.participants = participants
            duplicate_names = {
                name: rows for name, rows in name_rows.items() if len(rows) > 1
            }
            for name, duplicate_rows in duplicate_names.items():
                logger.warning(
                    "Duplicate participant name supported: name=%s rows=%s",
                    name,
                    duplicate_rows,
                )
            logger.info(
                "Participants loaded: count=%d ids=%s",
                len(participants),
                [participant["id"] for participant in participants],
            )
            return participants
        finally:
            workbook.close()

    @classmethod
    def _resolve_default_file(cls) -> Path:
        if cls.DEFAULT_REAL_FILE.exists():
            return cls.DEFAULT_REAL_FILE
        if cls.DEFAULT_EXAMPLE_FILE.exists():
            return cls.DEFAULT_EXAMPLE_FILE
        raise ParticipantDataError(
            "Nessun file partecipanti trovato. Inserire "
            "'data/partecipanti.xlsx' oppure "
            "'data/partecipanti_example.xlsx'."
        )
